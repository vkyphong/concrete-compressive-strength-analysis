"""Tạo Word trực tiếp từ Markdown chuẩn để hai bản luôn cùng nội dung.

Chạy: python analysis/build_report.py
Cần pandoc trong PATH (xem README). PDF được xuất từ chính DOCX bằng LibreOffice.
"""
from pathlib import Path
import shutil
import subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "report" / "hypothesis_tests_report.md"
WORD = ROOT / "report" / "Phan_3_Kiem_dinh_gia_thuyet.docx"


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if shutil.which("pandoc") is None:
        raise RuntimeError("Thiếu pandoc. Cài pandoc hoặc dùng Word đã cung cấp trong gói.")
    subprocess.run([
        "pandoc", str(SOURCE), "--from=markdown+tex_math_dollars",
        "--to=docx", f"--resource-path={SOURCE.parent}:{ROOT}",
        "--output", str(WORD),
    ], check=True)
    # Định dạng trên cùng DOCX, không thay nội dung hay thứ tự hình/bảng.
    doc = Document(WORD)
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.2)
    style_names = {style.name: style for style in doc.styles}
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        style = style_names.get(style_name)
        if style is None:
            continue
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    for style_name, size in (("Heading 1", 14), ("Heading 2", 12),
                             ("Heading 3", 11)):
        style = style_names.get(style_name)
        if style is not None:
            style.font.size = Pt(size)
            style.font.bold = True
    for table in doc.tables:
        # Tránh xé hàng và giữ toàn bộ bảng ngắn trong cùng một trang.
        for index, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            if properties.find(qn("w:cantSplit")) is None:
                properties.append(OxmlElement("w:cantSplit"))
            if index == 0 and properties.find(qn("w:tblHeader")) is None:
                properties.append(OxmlElement("w:tblHeader"))
            if index < len(table.rows) - 1:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Tài liệu tham khảo của phần 3"):
            paragraph.paragraph_format.page_break_before = True
    doc.save(WORD)
    print("Đã tạo Word từ cùng một nguồn Markdown:", WORD)


if __name__ == "__main__":
    main()
